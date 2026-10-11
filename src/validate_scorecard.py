#!/usr/bin/env python3
"""
validate_scorecard.py
=====================
Independent check of the Excel scorecard: re-implements the Power Query logic in pandas
and compares the result with the values cached in the workbook.

    python src/validate_scorecard.py \
        --csv data/raw/american_bankruptcy.csv \
        --xlsx excel/US_Bankruptcy_Risk_Scorecard.xlsx

Exit code 0 = every comparison matched (relative tolerance 1e-9).
"""
import argparse
import math
import sys

import numpy as np
import pandas as pd
from openpyxl import load_workbook

RENAME = {"X1": "Current_Assets", "X2": "COGS", "X3": "Depreciation_Amort", "X4": "EBITDA", "X5": "Inventory",
          "X6": "Net_Income", "X7": "Receivables", "X8": "Market_Value", "X9": "Net_Sales", "X10": "Total_Assets",
          "X11": "LT_Debt", "X12": "EBIT", "X13": "Gross_Profit", "X14": "Current_Liabilities",
          "X15": "Retained_Earnings", "X16": "Total_Revenue", "X17": "Total_Liabilities", "X18": "Operating_Expenses"}


def score(csv_path: str):
    """Re-implement the Power Query steps. Returns (all firm-years, one-row-per-company scores)."""
    df = pd.read_csv(csv_path, encoding="utf-8-sig").rename(columns=RENAME)
    df["Failed_Flag"] = (df["status_label"] == "failed").astype(int)
    df["Last_Year"] = df.groupby("company_name")["year"].transform("max")
    df["Is_Last_Year"] = (df["year"] == df["Last_Year"]).astype(int)
    df["Z_Valid"] = ((df["Total_Assets"] >= 1) & (df["Total_Liabilities"] >= 1)).astype(int)
    v = df["Z_Valid"] == 1
    df["WC_TA"] = np.where(v, (df["Current_Assets"] - df["Current_Liabilities"]) / df["Total_Assets"], np.nan)
    df["RE_TA"] = np.where(v, df["Retained_Earnings"] / df["Total_Assets"], np.nan)
    df["EBIT_TA"] = np.where(v, df["EBIT"] / df["Total_Assets"], np.nan)
    df["MVE_TL"] = np.where(v, df["Market_Value"] / df["Total_Liabilities"], np.nan)
    df["Sales_TA"] = np.where(v, df["Net_Sales"] / df["Total_Assets"], np.nan)
    df["Altman_Z"] = 1.2 * df.WC_TA + 1.4 * df.RE_TA + 3.3 * df.EBIT_TA + 0.6 * df.MVE_TL + 1.0 * df.Sales_TA
    df["Zone"] = np.where(~v, "Excluded", np.where(df.Altman_Z > 2.99, "Safe", np.where(df.Altman_Z >= 1.81, "Grey", "Distress")))
    c = df[(df.Is_Last_Year == 1) & v].sort_values("Altman_Z", kind="stable").reset_index(drop=True)
    c["Risk_Rank"] = np.arange(1, len(c) + 1)
    c["Risk_Decile"] = np.ceil(c.Risk_Rank * 10 / len(c)).astype(int)
    return df, c


def main(csv: str, xlsx: str) -> int:
    df, c = score(csv)
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    bad = 0

    def check(label, got, want):
        nonlocal bad
        ok = math.isclose(float(got), float(want), rel_tol=1e-9, abs_tol=1e-12)
        bad += (not ok)
        print(f"{'OK  ' if ok else 'FAIL'} {label:<44} python={got:<22} workbook={want}")

    heads = [x.value for x in next(wb["Dashboard"].iter_rows(min_row=5, max_row=5, min_col=2, max_col=6))]
    vals = [x.value for x in next(wb["Dashboard"].iter_rows(min_row=6, max_row=6, min_col=2, max_col=6))]
    d = dict(zip(heads, vals))
    n, f = len(c), int(c.Failed_Flag.sum())
    check("Companies analysed", n, d["Companies Analysed"])
    check("Bankruptcies", f, d["Bankruptcies"])
    check("Overall bankruptcy rate", f / n, d["Overall Bankruptcy Rate"])
    check("Bankruptcies flagged as Distress", c[(c.Zone == "Distress") & (c.Failed_Flag == 1)].shape[0] / f, d["Bankruptcies Flagged as Distress"])
    check("Captured in riskiest 30%", c[(c.Risk_Decile <= 3) & (c.Failed_Flag == 1)].shape[0] / f, d["Captured in Riskiest 30%"])

    zv = {r[0].value: [x.value for x in r] for r in wb["Zone Validation"].iter_rows(min_row=5, max_row=7, max_col=5)}
    for z in ("Distress", "Grey", "Safe"):
        g = c[c.Zone == z]
        check(f"Zone {z}: companies", len(g), zv[z][1])
        check(f"Zone {z}: bankrupt", int(g.Failed_Flag.sum()), zv[z][2])

    dl = {int(r[0].value): [x.value for x in r] for r in wb["Decile Lift"].iter_rows(min_row=4, max_row=13, max_col=6)}
    cum = 0
    for k in range(1, 11):
        g = c[c.Risk_Decile == k]
        cum += int(g.Failed_Flag.sum())
        check(f"Decile {k}: companies", len(g), dl[k][1])
        check(f"Decile {k}: bankrupt", int(g.Failed_Flag.sum()), dl[k][2])
        check(f"Decile {k}: cumulative capture", cum / f, dl[k][5])

    tr = {int(r[0].value): r[1].value for r in wb["Bankruptcy Trend"].iter_rows(min_row=4, max_row=23, max_col=2)}
    by_year = c[c.Failed_Flag == 1].groupby("year").size()
    for y in range(1999, 2019):
        check(f"Bankruptcies in {y}", int(by_year.get(y, 0)), tr[y])

    rc = {r[0].value: [x.value for x in r] for r in wb["Ratio Comparison"].iter_rows(min_row=4, max_row=9, max_col=3)}
    for label, col in [("Working Capital / Assets", "WC_TA"), ("Retained Earnings / Assets", "RE_TA"), ("EBIT / Assets", "EBIT_TA"),
                       ("Market Value / Liabilities", "MVE_TL"), ("Sales / Assets", "Sales_TA"), ("Altman Z-Score", "Altman_Z")]:
        check(f"Median {col}: healthy", c[c.status_label == "alive"][col].median(), rc[label][1])
        check(f"Median {col}: failed", c[c.status_label == "failed"][col].median(), rc[label][2])

    print(f"\n{'ALL CHECKS PASSED' if not bad else str(bad) + ' CHECK(S) FAILED'}")
    return 1 if bad else 0


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--csv", default="data/raw/american_bankruptcy.csv")
    ap.add_argument("--xlsx", default="excel/US_Bankruptcy_Risk_Scorecard.xlsx")
    a = ap.parse_args()
    sys.exit(main(a.csv, a.xlsx))
